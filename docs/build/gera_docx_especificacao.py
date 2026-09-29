"""DOCX da especificação do projeto para a equipe, a partir de docs/especificacao_projeto.md (pandoc + ajustes).

    python docs/build/diagrama_png.py <pasta_tmp>/diagrama.png      # diagrama Mermaid -> PNG (Chrome + CDN do Mermaid)
    python docs/build/gera_docx_especificacao.py <pasta_tmp> docs/especificacao_projeto.docx

Requer pandoc, python-docx e Playwright (requirements-dev). Rodar da raiz, depois de editar o .md."""
import re, subprocess, sys, pathlib
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

S = pathlib.Path(sys.argv[1]); SAIDA = pathlib.Path(sys.argv[2])
AZUL = RGBColor(0x00, 0x4A, 0x80)

# 1. markdown para o Word: sem o sumário manual (o Word gera o dele), diagrama como imagem, sem links internos
md = pathlib.Path('docs/especificacao_projeto.md').read_text(encoding='utf-8')
md = re.sub(r'^# .*\n', '', md, count=1)                                   # o título vai na capa
md = re.sub(r'## Sumário\n.*?\n---\n', '', md, count=1, flags=re.S)
md = re.sub(r'```mermaid\n.*?```', f'![Fluxo de dados: das fontes ao site, ao PDF e ao DOCX de curadoria]({(S / "diagrama.png").as_posix()}){{width=15.5cm}}', md, flags=re.S)
md = re.sub(r'\[([^\]]+)\]\(#[^)]+\)', r'\1', md)
md = re.sub(r'^---\s*$', '', md, flags=re.M)                               # réguas horizontais: a quebra vem dos títulos
(S / 'spec.md').write_text(md, encoding='utf-8')

# 2. documento de referência (estilos)
ref = S / 'ref.docx'
ref.write_bytes(subprocess.run(['pandoc', '--print-default-data-file', 'reference.docx'], capture_output=True).stdout)
d = Document(ref)
st = d.styles
def fonte(nome, tam=None, cor=None, negrito=None, familia='Calibri'):
    s = st[nome]; f = s.font; f.name = familia
    s.element.rPr.rFonts.set(qn('w:eastAsia'), familia) if s.element.rPr is not None and s.element.rPr.rFonts is not None else None
    if tam: f.size = Pt(tam)
    if cor is not None: f.color.rgb = cor
    if negrito is not None: f.bold = negrito
for n in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
    if n in [x.name for x in st]:
        fonte(n, 10.5)
for n, t in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)]:
    fonte(n, t, AZUL, True)
for n, t in [('Title', 24), ('Subtitle', 13)]:
    fonte(n, t, AZUL, n == 'Title')
fonte('Date', 10.5, RGBColor(0x55, 0x5F, 0x6D))
if 'Block Text' in [x.name for x in st]:
    fonte('Block Text', 10, RGBColor(0x33, 0x3F, 0x4D))
for n in ['Source Code', 'Verbatim Char']:
    if n in [x.name for x in st]:
        st[n].font.name = 'Consolas'; st[n].font.size = Pt(9)
sec = d.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.2)
d.save(ref)

# 3. pandoc
subprocess.run(['pandoc', str(S / 'spec.md'), '-f', 'gfm+attributes+implicit_figures', '-o', str(SAIDA), '--reference-doc', str(ref), '--toc',
                '--toc-depth=2', '--shift-heading-level-by=-1', '-M', 'title=Diagnóstico da Primeira Infância Carioca',
                '-M', 'subtitle=Especificação funcional e técnica do projeto',
                '-M', 'date=Instituto Pereira Passos · versão de 28 de setembro de 2026',
                '-M', 'toc-title=Sumário', '--resource-path', str(S)], check=True)

# 4. ajustes: bordas e cabeçalho nas tabelas, largura total, quebra de página depois do sumário
d = Document(SAIDA)
def borda(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for lado in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{lado}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'B7C4D1')
        b.append(e)
    tblPr.append(b)
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW'); tblPr.append(w)
    w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), '5000')
for tbl in d.tables:
    borda(tbl)
    for i, row in enumerate(tbl.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i == 0:
                        r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
                sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), '004A80'); tcPr.append(sh)
    # cabeçalho repete em cada página
    trPr = tbl.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
    # largura das colunas pelo conteúdo (o pandoc usa o tamanho dos traços do markdown, que esprema texto longo)
    LARG = 16.0   # cm úteis
    pesos = []
    for j in range(len(tbl.columns)):
        textos = [len(row.cells[j].text) for row in tbl.rows if j < len(row.cells)]
        media = sum(textos) / len(textos)
        pesos.append(max(3, min(60, media * 0.8 + max(textos) * 0.2)))
    # mínimo: a maior palavra da coluna cabe inteira (código como `ibge_censo2022` não quebra no meio)
    minimos = [max(1.2, 0.2 * max(len(t) for row in tbl.rows[1:] if j < len(row.cells) for t in (row.cells[j].text.split() or [''])) + 0.3)
               for j in range(len(tbl.columns))]
    tot = sum(pesos); cms = [max(m, LARG * p / tot) for p, m in zip(pesos, minimos)]
    excesso = sum(cms) - LARG
    if excesso > 0:   # tira da(s) coluna(s) com folga acima do mínimo, proporcionalmente
        folga = [c - m for c, m in zip(cms, minimos)]; f = sum(folga)
        cms = [c - excesso * fo / f for c, fo in zip(cms, folga)] if f > 0 else [c * LARG / sum(cms) for c in cms]
    grid = tbl._tbl.tblGrid
    for gc, c in zip(grid.findall(qn('w:gridCol')), cms):
        gc.set(qn('w:w'), str(int(c * 567)))
    for row in tbl.rows:
        for cell, c in zip(row.cells, cms):
            cell.width = Cm(c)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tbl._tbl.tblPr.append(lay)
    w = tbl._tbl.tblPr.find(qn('w:tblW')); w.set(qn('w:type'), 'dxa'); w.set(qn('w:w'), str(int(LARG * 567)))
# quebra de página só depois do sumário (antes do capítulo 1)
for p in d.paragraphs:
    if p.style.name == 'Heading 1':
        p.paragraph_format.page_break_before = True
        break
# sumário estático (capítulos e subseções): o campo TOC do pandoc só aparece depois de "atualizar campos" no Word
cabecalhos = [(p.style.name, p.text) for p in d.paragraphs if p.style.name in ('Heading 1', 'Heading 2')]
toc = next(el for el in d.element.body.iter(qn('w:sdt')))   # bloco do sumário gerado pelo pandoc
novo = []
tit = OxmlElement('w:p'); d.element.body.insert(list(d.element.body).index(toc), tit)
from docx.text.paragraph import Paragraph
pt = Paragraph(tit, d._body); r = pt.add_run('Sumário'); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL
ancora = tit
for estilo, texto in cabecalhos:
    el = OxmlElement('w:p'); ancora.addnext(el); ancora = el
    p = Paragraph(el, d._body); run = p.add_run(texto); run.font.size = Pt(10 if estilo == 'Heading 1' else 9.5)
    run.bold = estilo == 'Heading 1'
    p.paragraph_format.left_indent = Cm(0 if estilo == 'Heading 1' else 0.8)
    p.paragraph_format.space_after = Pt(1 if estilo == 'Heading 2' else 2)
toc.getparent().remove(toc)
# rodapé com numeração de página
rod = d.sections[0].footer.paragraphs[0]; rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = rod.add_run('Diagnóstico da Primeira Infância Carioca · especificação do projeto · página ')
r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x55, 0x5F, 0x6D)
r2 = rod.add_run(); r2.font.size = Pt(8)
for tag, txt in [('begin', None), (None, 'PAGE'), ('end', None)]:
    if tag:
        f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), tag); r2._r.append(f)
    else:
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r2._r.append(it)
d.save(SAIDA)
print('ok', SAIDA, SAIDA.stat().st_size)
